"""Evidence must survive the Word-to-PDF extraction intermediate."""
import io

import fitz
from docx import Document

from app.extraction.docx_ingest import docx_to_pdf_bytes


def _convert(document):
    stream = io.BytesIO()
    document.save(stream)
    return fitz.open(stream=docx_to_pdf_bytes(stream.getvalue()), filetype="pdf")


def test_wide_table_preserves_rightmost_identity_and_amount():
    document = Document()
    row = document.add_table(rows=1, cols=5).rows[0]
    values = ["Long repair description " * 12, "Registration", "XY26ABC", "Claim: CLAIM-98765", "1234.56"]
    for cell, value in zip(row.cells, values, strict=True):
        cell.text = value
    with _convert(document) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        for value in values[1:]:
            assert value in text
        assert pdf[0].rect.width > 595


def test_header_and_footer_identity_are_preserved():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Claim number: HEADER-123"
    document.add_paragraph("Engineer assessment body")
    document.sections[0].footer.paragraphs[0].text = "Policy number: FOOTER-456"
    with _convert(document) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        assert "Claim number: HEADER-123" in text
        assert "Policy number: FOOTER-456" in text
        assert "Engineer assessment body" in text
