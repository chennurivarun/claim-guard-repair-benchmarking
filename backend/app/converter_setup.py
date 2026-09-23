"""Set up and verify the local Word-to-PDF converter; never process client files."""

from __future__ import annotations

import argparse
import io
import shutil
import subprocess
import sys

import fitz
from docx import Document
from PIL import Image

from app.services.document_processing import find_libreoffice, normalise_document_upload


def build_conversion_probe() -> bytes:
    """Exercise images, headers and nested tables that need the Office fallback."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "CONVERTER HEADER 923"
    document.add_paragraph("CONVERTER INVOICE 923")
    image = io.BytesIO()
    Image.new("RGB", (32, 32), (245, 88, 0)).save(image, format="PNG")
    document.add_picture(io.BytesIO(image.getvalue()))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "CONVERTER TABLE 923"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "CONVERTER NESTED 923"
    content = io.BytesIO()
    document.save(content)
    return content.getvalue()


def verify_conversion() -> int:
    result = normalise_document_upload("converter-check.docx", build_conversion_probe())
    with fitz.open(stream=result.content, filetype="pdf") as pdf:
        text = " ".join(" ".join(page.get_text().split()) for page in pdf)
        markers = ("HEADER", "INVOICE", "TABLE", "NESTED")
        if not all(f"CONVERTER {marker} 923" in text for marker in markers):
            raise ValueError("The test PDF is missing text from its header or tables.")
        if not any(page.get_images() for page in pdf):
            raise ValueError("The test PDF is missing its image.")
        return len(pdf)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--install", action="store_true",
        help="If missing on Windows, install LibreOffice through WinGet, then test conversion.",
    )
    args = parser.parse_args(argv)
    print("Word conversion setup (synthetic test only; no AI calls or database changes)")
    try:
        executable = find_libreoffice()
    except ValueError as exc:
        print(f"FAIL: {exc}")
        return 1

    if not executable and args.install:
        if sys.platform != "win32":
            print("Automatic installation is supported on Windows. Install LibreOffice "
                  "using your operating system's installer, then run this check again.")
            return 1
        winget = shutil.which("winget")
        if not winget:
            print("WinGet is unavailable. Ask IT to install LibreOffice on the backend "
                  "computer, then rerun uv run claimguard-converter.")
            return 1
        print("Installing LibreOffice via WinGet. Follow any installer or organisation "
              "approval prompts. This may take several minutes.", flush=True)
        try:
            # Inherit the terminal: agreements and administrator prompts remain visible.
            installed = subprocess.run(
                [winget, "install", "--id", "TheDocumentFoundation.LibreOffice",
                 "--exact", "--source", "winget"],
                check=False, timeout=900,
            )
        except (OSError, subprocess.SubprocessError):
            print("FAIL: The installer could not finish. Check the installer output or "
                  "ask IT to install LibreOffice, then rerun this check.")
            return 1
        if installed.returncode != 0:
            print(f"FAIL: WinGet returned {installed.returncode}. Resolve the installation "
                  "issue shown above, then rerun this command.")
            return 1
        executable = find_libreoffice()

    if not executable:
        print("FAIL: LibreOffice was not found. On Windows run: "
              "uv run claimguard-converter --install")
        print("For a custom installation, set CLAIM_GUARD_LIBREOFFICE_PATH to the full "
              "executable path in backend/.env. Otherwise install LibreOffice on this computer.")
        return 1
    print(f"Converter: {executable}")
    try:
        pages = verify_conversion()
    except Exception as exc:
        detail = str(exc) if isinstance(exc, ValueError) else type(exc).__name__
        print(f"FAIL: {detail}")
        return 1
    print(f"PASS: Word file with an image, header and nested table converted to {pages} PDF page(s).")
    print("Complex Word uploads will use this converter automatically; originals stay unchanged.")
    print("Restart the backend, then upload previously rejected files again. "
          "This checks conversion only; scanned pages still require OCR or image extraction.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
